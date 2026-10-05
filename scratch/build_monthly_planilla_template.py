#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera la Plantilla Oficial Mensual de Médicos en Planilla - DoctorSV
Contiene 2 hojas:
  - 'Edward Zelaya' (Grupo 1)
  - 'Roxana Canales' (Grupo 2)
Con 31 días (01/10/2026 a 31/10/2026) y sólo los campos requeridos por cada supervisor.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from datetime import datetime
import os
import shutil

# IDs corregidos y limpios
KNOWN_IDS = {
    'EMERITA SARAI CRUZ CANIZALEZ': '052501581',
    'HEIDI JULISSA BURUCA RIVERA': '043865009',
    'HEYDI JULISSA BURUCA RIVERA': '043865009',
    'RICARDO DANIEL CHICAS': '054336429',
    'RICARDO DANIEL CHICAS CHICAS': '054336429',
    'SANTOS BLADIMIR CHAVEZ CASTILLO': '052938754',
    'ZULEYMA ABIGAIL GARCIA REYES': '061763310',
}

def load_source_data():
    wb_src = openpyxl.load_workbook('PlanillaMensual.xlsx', data_only=True)
    ws_src = wb_src['SEDE SAN MIGUEL']

    # 1. Extraer Edward Zelaya y su grupo
    edward_data = {
        'supervisor_name': 'EDWARD JOSUE ZELAYA PRUDENCIO',
        'supervisor_role': 'SUPERVISOR',
        'supervisor_id': '049903944',
        'grupo': 'GRUPO 1',
        'puesto': '139',
        'supervisor_shifts': [ws_src.cell(3, c).value or 'LIBRE' for c in range(7, 38)],
        'doctors': []
    }
    for r in range(4, 32):
        name = ws_src.cell(r, 4).value
        func = ws_src.cell(r, 5).value or 'CONSULTANTE'
        doc_id = ws_src.cell(r, 6).value
        if not name:
            continue
        clean_name = str(name).strip()
        clean_id = KNOWN_IDS.get(clean_name, str(doc_id).strip() if doc_id else '')
        if clean_id in ['VALIDAR NOMBRE', '#N/A', 'None']:
            clean_id = KNOWN_IDS.get(clean_name, '')
        shifts = [ws_src.cell(r, c).value or 'LIBRE' for c in range(7, 38)]
        edward_data['doctors'].append({
            'num': len(edward_data['doctors']) + 1,
            'name': clean_name,
            'role': str(func).strip(),
            'id': clean_id,
            'shifts': shifts
        })

    # 2. Extraer Roxana Canales y su grupo
    roxana_data = {
        'supervisor_name': 'ROXANA GUADALUPE CANALES RODRIGUEZ',
        'supervisor_role': 'SUPERVISORA',
        'supervisor_id': '056774097',
        'grupo': 'GRUPO 2',
        'puesto': '138',
        'supervisor_shifts': [ws_src.cell(33, c).value or 'LIBRE' for c in range(7, 38)],
        'doctors': []
    }
    for r in range(34, 63):
        name = ws_src.cell(r, 4).value
        func = ws_src.cell(r, 5).value or 'CONSULTANTE'
        doc_id = ws_src.cell(r, 6).value
        if not name:
            continue
        clean_name = str(name).strip()
        clean_id = KNOWN_IDS.get(clean_name, str(doc_id).strip() if doc_id else '')
        if clean_id in ['VALIDAR NOMBRE', '#N/A', 'None']:
            clean_id = KNOWN_IDS.get(clean_name, '')
        shifts = [ws_src.cell(r, c).value or 'LIBRE' for c in range(7, 38)]
        roxana_data['doctors'].append({
            'num': len(roxana_data['doctors']) + 1,
            'name': clean_name,
            'role': str(func).strip(),
            'id': clean_id,
            'shifts': shifts
        })

    return edward_data, roxana_data

def build_supervisor_sheet(ws, data):
    # Tipografía moderna
    font_main_title = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    font_sub_title = Font(name="Segoe UI", size=10.5, bold=True, color="E2E8F0")
    font_hdr = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_sup = Font(name="Segoe UI", size=10.5, bold=True, color="003487")
    font_doc = Font(name="Segoe UI", size=9.5, color="1E293B")
    font_doc_bold = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
    font_center = Font(name="Segoe UI", size=9.5, color="334155")
    font_shift = Font(name="Segoe UI", size=9, bold=True, color="1E293B")

    # Colores corporativos DoctorSV
    fill_banner = PatternFill("solid", fgColor="0048B5")     # Azul institucional
    fill_banner_sub = PatternFill("solid", fgColor="0F3B82") # Azul oscuro elegante
    fill_hdr = PatternFill("solid", fgColor="0048B5")        # Azul encabezados
    fill_sup = PatternFill("solid", fgColor="EBF3FF")        # Resaltado supervisor
    fill_zebra_1 = PatternFill("solid", fgColor="FFFFFF")    # Blanco
    fill_zebra_2 = PatternFill("solid", fgColor="F8FAFC")    # Slate 50 muy tenue

    border_thin = Side(style="thin", color="CBD5E1")
    border_double = Side(style="double", color="94A3B8")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    sup_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_double)

    # 1. Fila 1: Banner Principal
    ws.merge_cells("A1:AI1")
    ws["A1"] = f"DOCTORSV · PROGRAMACIÓN MENSUAL DE MÉDICOS EN PLANILLA — {data['grupo'].upper()}"
    ws["A1"].font = font_main_title
    ws["A1"].fill = fill_banner
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # 2. Fila 2: Sub-banner informativo
    ws.merge_cells("A2:AI2")
    ws["A2"] = (
        f"SUPERVISOR: {data['supervisor_name']}  |  "
        f"ESTACIÓN PUESTO #{data['puesto']}  |  "
        f"MES: OCTUBRE 2026 (31 DÍAS)  |  "
        f"TOTAL MÉDICOS: {len(data['doctors'])} CONSULTANTES"
    )
    ws["A2"].font = font_sub_title
    ws["A2"].fill = fill_banner_sub
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # 3. Fila 3: Espaciador
    ws.row_dimensions[3].height = 8

    # 4. Fila 4: Encabezados de tabla
    headers = ["N°", "MÉDICO / COLABORADOR", "FUNCIÓN", "IDENTIFICADOR"]
    for day in range(1, 32):
        headers.append(f"{day:02d}/10/2026")

    ws.row_dimensions[4].height = 26
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=c_idx, value=h)
        cell.font = font_hdr
        cell.fill = fill_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border

    # Anchos de columna
    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    for c in range(5, 36):
        col_letter = get_column_letter(c)
        ws.column_dimensions[col_letter].width = 17

    # 5. Fila 5: Fila del Supervisor (Destacada)
    ws.row_dimensions[5].height = 24
    row5_vals = ["SUP", data['supervisor_name'], data['supervisor_role'], data['supervisor_id']] + data['supervisor_shifts']
    for c_idx, val in enumerate(row5_vals, 1):
        clean_val = str(val).strip() if val else 'LIBRE'
        cell = ws.cell(row=5, column=c_idx, value=clean_val)
        cell.fill = fill_sup
        cell.border = sup_border
        if c_idx == 1:
            cell.font = font_sup
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif c_idx == 2:
            cell.font = font_sup
            cell.alignment = Alignment(horizontal="left", vertical="center")
        elif c_idx in (3, 4):
            cell.font = font_sup
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.font = font_shift
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # 6. Filas 6+: Médicos consultantes
    for idx, doc in enumerate(data['doctors'], 1):
        r_num = 5 + idx
        ws.row_dimensions[r_num].height = 21
        row_fill = fill_zebra_1 if idx % 2 == 1 else fill_zebra_2

        row_vals = [doc['num'], doc['name'], doc['role'], doc['id']] + doc['shifts']
        for c_idx, val in enumerate(row_vals, 1):
            clean_val = str(val).strip() if val else 'LIBRE'
            cell = ws.cell(row=r_num, column=c_idx, value=clean_val)
            cell.fill = row_fill
            cell.border = cell_border
            if c_idx == 1:
                cell.font = font_center
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 2:
                cell.font = font_doc_bold
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c_idx in (3, 4):
                cell.font = font_center
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.font = font_shift
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # 7. Regla de validación de datos para turnos (Menú desplegable en celdas de fechas)
    dv = DataValidation(
        type="list",
        formula1='"06:00am-02:00pm,07:00am-03:00pm,02:00pm-10:00pm,10:00pm-06:00am,LIBRE,VACACIÓN,INCAPACIDAD,MATERNIDAD,PERMISO PERSONAL,PERMUTA"',
        allow_blank=True
    )
    dv.error = "Por favor selecciona un turno oficial válido de la lista."
    dv.errorTitle = "Turno No Válido"
    dv.prompt = "Selecciona el turno oficial o estado del médico."
    dv.promptTitle = "Turno Mensual"
    ws.add_data_validation(dv)
    
    last_row = 5 + len(data['doctors'])
    dv.add(f"E5:AI{last_row}")

    # 8. Inmovilizar paneles en E5 (columnas fijas A-D y encabezados 1-4)
    ws.freeze_panes = "E5"

def generate_monthly_template():
    edward_data, roxana_data = load_source_data()

    wb = openpyxl.Workbook()
    
    # Hoja 1: Edward Zelaya
    ws_edward = wb.active
    ws_edward.title = "Edward Zelaya"
    build_supervisor_sheet(ws_edward, edward_data)

    # Hoja 2: Roxana Canales
    ws_roxana = wb.create_sheet("Roxana Canales")
    build_supervisor_sheet(ws_roxana, roxana_data)

    output_root = "Plantilla_Planilla_Mensual_DoctorSV.xlsx"
    wb.save(output_root)
    print(f"✅ Archivo principal creado: {output_root}")

    # Copiar a public y dist para descarga directa desde la app web
    public_target = "public/Plantilla_Planilla_Mensual_DoctorSV.xlsx"
    shutil.copyfile(output_root, public_target)
    print(f"✅ Copiado a public: {public_target}")

    # También actualizar la plantilla general para mantener compatibilidad
    shutil.copyfile(output_root, "public/Plantilla_Planilla_DoctorSV.xlsx")
    if os.path.exists("dist"):
        shutil.copyfile(output_root, "dist/Plantilla_Planilla_Mensual_DoctorSV.xlsx")
        shutil.copyfile(output_root, "dist/Plantilla_Planilla_DoctorSV.xlsx")
        print("✅ Copiado a dist")

if __name__ == "__main__":
    generate_monthly_template()
